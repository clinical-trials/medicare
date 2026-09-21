"""Rebuild the journal-style manuscript and supplements from auditable sources.

Run with the bundled Codex Python runtime. Render with the packaged DOCX renderer;
the PDF is copied from that renderer only after visual review of every page.
"""
from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
OUT=ROOT/'outputs/01a0ab2c-c289-74e3-8442-d7966be4cbe8'
STEM='Medicare_equity_research_and_health_IT_manuscript_2026-09-21'
read=lambda path:json.loads(path.read_text())
source=(HERE/'manuscript.md').read_text()
opportunities=read(HERE/'opportunities.json')
tables=read(HERE/'publication_tables.json')
registry=read(ROOT/'research/reference_registry.json')
supp=(HERE/'supplemental_methods.md').read_text()
refs={r['id']:r for k in ['scientific_references','methodology_references','retrieval_references','translational_references','policy_sources','clinical_guidance_sources'] for r in registry[k]}
used=[]
def note_refs(text):
    for rid in re.findall(r'\b(?:L-[A-Z]+\d+(?:-[A-Z]\d+)?|[MPC]-\d{3})\b',text):
        if rid not in used:used.append(rid)

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72)
sec.left_margin=sec.right_margin=Inches(.78)
sec.header_distance=sec.footer_distance=Inches(.30)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Bullet']:
    st=styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.font._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    st.paragraph_format.widow_control=True
    rpr=st.element.find(qn('w:rPr'))
    if rpr is not None:
        color=rpr.find(qn('w:color'))
        if color is not None:
            for a in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+a),None)
    ppr=st.element.find(qn('w:pPr'))
    if ppr is not None:
        for elem in list(ppr):
            if elem.tag in [qn('w:pBdr'),qn('w:shd')]:ppr.remove(elem)
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.line_spacing=1.12
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Title'].font.size=Pt(20);styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(8)
styles['Subtitle'].font.size=Pt(12);styles['Subtitle'].font.italic=False
for name,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
    styles[name].font.size=Pt(size);styles[name].font.bold=True
    styles[name].paragraph_format.space_before=Pt(12)
    styles[name].paragraph_format.space_after=Pt(5)
    styles[name].paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(10)
styles['Caption'].font.italic=False
styles['Caption'].font.bold=False
styles['Caption'].paragraph_format.space_after=Pt(7)
styles['List Bullet'].paragraph_format.space_after=Pt(5)
styles['List Bullet'].paragraph_format.line_spacing=1.06

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
footer.style=styles['Normal'];footer.paragraph_format.space_after=Pt(0)

def text_runs(p,text):
    note_refs(text)
    for t in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`|\*[^*]+\*)',text):
        if t.startswith('**') and t.endswith('**'):
            p.add_run(t[2:-2]).bold=True
        elif t.startswith('[') and '](' in t:
            label,url=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',t).groups()
            if not url.startswith('http'):
                target=(HERE/url).resolve().relative_to(ROOT)
                url='https://github.com/clinical-trials/medicare/blob/main/'+str(target)
            hyperlink(p,label,url,11)
        elif t.startswith('`') and t.endswith('`'):
            p.add_run(t[1:-1])
        elif t.startswith('*') and t.endswith('*'):
            p.add_run(t[1:-1]).italic=True
        else:p.add_run(t)
    return p

def para(text,style=None):
    return text_runs(doc.add_paragraph(style=style),text)

pending_page=False
def heading(text,level=1):
    global pending_page
    p=doc.add_heading(re.sub(r'\s+',' ',re.sub(r'[^\w\s]',' ',text)).strip(),level)
    if pending_page:p.paragraph_format.page_break_before=True;pending_page=False
    return p

def page():
    global pending_page
    pending_page=True

def figure_image(path,description):
    global pending_page
    p=doc.add_paragraph()
    if pending_page:p.paragraph_format.page_break_before=True;pending_page=False
    shape=p.add_run().add_picture(str(path),width=Inches(6.85))
    shape._inline.docPr.set('descr',description)

def markdown(text,skip_title=False):
    # Paragraph-level Markdown is deliberately restricted for a reproducible manuscript.
    for block in re.split(r'\n\s*\n',text.strip()):
        lines=block.splitlines()
        if lines[0].startswith('#'):
            if skip_title and lines[0].startswith('# '):continue
            n=len(lines[0])-len(lines[0].lstrip('#'))
            heading(lines[0][n:].strip(),max(1,min(n-1,3)))
        elif all(line.startswith('- ') for line in lines):
            for line in lines:
                value=line[2:].replace('[x] ','Completed: ').replace('[ ] ','Pending: ')
                para(value,'List Bullet')
        elif all(re.match(r'\d+\. ',line) for line in lines):
            for line in lines:para(line)
        elif lines[0].startswith('|'):
            cells=[[c.strip() for c in line.strip('|').split('|')] for line in lines]
            table({'headers':cells[0],'rows':cells[2:],'note':''},[2.25,4.65])
        else:
            value=' '.join(lines)
            para(value)

def hyperlink(p,label,url,size=10):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'173F5F');pr.append(color)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));pr.append(sz)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def table(data,widths):
    t=doc.add_table(rows=1,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,w,label in zip(t.rows[0].cells,widths,data['headers']):
        c.width=Inches(w);text_runs(c.paragraphs[0],label)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in data['rows']:
        cells=t.add_row().cells
        for c,w,value in zip(cells,widths,row):
            c.width=Inches(w);text_runs(c.paragraphs[0],value)
    for idx,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            cp.append(borders)
            mar=OxmlElement('w:tcMar')
            for side,twips in [('top',85),('bottom',85),('left',95),('right',95)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(twips));e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'173F5F' if idx==0 else 'F2F5F7' if idx%2==0 else 'FFFFFF');cp.append(shd)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.line_spacing=1.04;p.paragraph_format.widow_control=False
                for r in p.runs:
                    r.font.size=Pt(9.5);r.font.name='Times New Roman'
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if idx==0 else '000000');r.bold=(idx==0)
    if data['note']:para(data['note'],'Caption')

title=source.splitlines()[0][2:]
doc.add_paragraph(title,'Title')
doc.add_paragraph('Evidence synthesis and methods for a systematic review in progress','Subtitle')
para('September 21 2026')
first=source.split('## Abstract',1)[1].split('## Introduction',1)[0]
markdown('## Abstract'+first)
para('**Keywords:** Medicare; health equity; sex differences; reimbursement; health information technology; evidence synthesis.')
page()
markdown('## Introduction'+source.split('## Introduction',1)[1])

page();figure_image(HERE/'figures/figure1_retrieval_status.png','Figure 1 Retrieval and review status. 59,147 unique report records identified; human screening and appraisal pending. Selected 38-record register shown separately.')
para('Counts describe executed PubMed development searches as of September 21, 2026. Each unique PMID is a report record, not necessarily an independent study. Search extensions were developed after the initial search. No counts of full-text exclusions or final included studies are inferred. The selected 38-record register and targeted intervention retrieval are separate routes requiring later reconciliation. This is an audit diagram, not a completed PRISMA selection flow.','Caption')

page();heading(tables['table1']['title']);table(tables['table1'],[1.32,2.68,2.90])
page();figure_image(HERE/'figures/figure2_validation_pathway.png','Figure 2 Health IT validation pathway. Need, mechanism, intervention, benefit and equity, and implementation require distinct evidence.')
para('Conceptual framework developed for this project after exploratory evidence review. Each stage requires a specified patient population, clinical indication and appropriate comparator. Progress may reveal that clinical support, policy change or an existing service is more useful than new software. Average benefit, equitable reach and commercial feasibility require separate evidence; no composite score or profitability ranking is implied.','Caption')
page();heading(tables['table2']['title']);table(tables['table2'],[1.22,2.7,2.98])

page();heading('Supplement A Methods and review completion')
markdown(supp,skip_title=True)

# Exact executed queries, with stable filenames for complete translated responses.
heading('Executed search sources',2)
para('The complete original and corrected core queries, date partitions, returned translations, warnings, ID memberships and response hashes are retained under research/review_2026-09-21/pubmed. Targeted extension queries and raw responses are retained under research/publication_2026-09-21/search_extensions. The consolidated queue retains all 59,147 records without automatic exclusions. The following extension strings reproduce the submitted PubMed syntax.')
for q in read(HERE/'search_extensions/queries.json')['queries']:
    para('**'+q['id']+'** '+q['purpose'])
    p=doc.add_paragraph(q['query']);p.paragraph_format.line_spacing=1.0
    for r in p.runs:r.font.name='Courier New';r.font.size=Pt(9)

page();heading('Supplement B Opportunity assessment details')
para('These six domains are research hypotheses without a comparative investment ranking. The accompanying workbook preserves the full evidence, workflow, source and policy fields. Prospective users and buyers have not been interviewed or validated. Proposed trials are future studies, not registrations or commitments. The patient population and separate measurement of sex and gender identity require an explicit protocol.')
for index,d in enumerate(opportunities['domains']):
    if index:page()
    heading('Opportunity '+str(index+1)+' '+d['domain'],2)
    para('**'+d['id']+'** '+d['problem'])
    para('**Clinical basis:** '+d['clinical_basis'])
    para('**Intervention to test:** '+d['proposed_intervention']['type']+'.')
    para('**Patient benefit:** '+d['patient_benefit_outcomes'][0])
    para('**Evidence anchors:** '+ '; '.join(e['study_id']+' ('+e['role']+')' for e in d['evidence'])+'.')
    para('**Prospective users:** '+ '; '.join(d['users'])+'. **Prospective buyers, unverified:** '+ '; '.join(d['potential_buyers'])+'.')
    para('**Data requirement:** '+d['data_and_workflow_needs'][0])
    para('**Policy boundary:** '+d['policy_boundary']+' ['+'; '.join(d['policy_source_ids'])+']')
    para('**Validation design:** '+d['validation_study']['design'])
    para('**Analysis:** '+d['validation_study']['analysis'])
    para('**Success criterion:** '+d['validation_study']['success_rule'])
    para('**Important risk:** '+d['alternatives_and_risks'][0])
    para('**Unresolved evidence:** '+d['evidence_gap'])

page();heading('Supplement C Reference roles and reporting declarations')
roles=read(HERE/'evidence_roles.json')
para('The legacy register contains 31 broadly labeled candidate records and seven contextual records. These labels are not final eligibility decisions. A separate publication-role audit identifies the following provisional uses; directness does not resolve patient identity or methodological validity.')
labels={
 'medicare_patient_empirical_candidate':'Potential Medicare patient empirical evidence',
 'medicare_descriptive_drug_use':'Medicare descriptive drug use and spending',
 'medicare_service_valuation_candidate':'Potential Medicare service valuation evidence',
 'valuation_input_directness_unresolved':'Service valuation inputs with unresolved Medicare directness',
 'clinical_efficacy_safety_background':'Clinical efficacy or safety background',
 'other_payer_empirical_background':'Other or mixed payer background',
 'physician_gender_context':'Physician gender context',
 'general_or_qualitative_context':'General or qualitative context'}
for k,label in labels.items():
    ids=[r['study_id'] for r in roles['records'] if r['publication_role_proposed']==k]
    # IDs in the role audit are inventory identifiers, not additional substantive citations.
    p=doc.add_paragraph();p.add_run(f'{label}: {len(ids)} records. ').bold=True;p.add_run(', '.join(ids)+'.')
para('Three newly examined intervention study families are retained separately as translational context. L-T03 and L-T03-P1 are reports of the same trial. The bibliography preserves excluded records and corrections; a correction is not an independent clinical study.')
heading('Authorship funding and interests',2)
para('Author names, affiliations, corresponding author, contributions, funding and competing interests have not been supplied or approved. No absence of conflicts or funding is asserted. Responsible investigators must review the complete manuscript, verify the scientific judgments and approve any submission. No journal inquiry or submission was sent in preparing this package.')
heading('Use of artificial intelligence',2)
para('OpenAI Codex assisted with search development and execution, public-source retrieval, preliminary screening notes and data extraction, bibliographic checking, code, evidence organization, prose and figures. Parallel AI checks were used for arithmetic, citation consistency and document review. These checks do not constitute two independent human reviewers or completed formal appraisal. Investigators retain responsibility for verification, eligibility decisions, interpretation and the final manuscript. The disclosure should be updated to reflect actual human contributions and the submission journal’s requirements.')
heading('Data code and access',2)
para('The repository retains queries, successful raw response captures, identifiers, hashes, provisional extractions, source locations, exclusions, bibliography, RIS export and artifact builders. Public-source availability varies; abstract-only and targeted full-text access are identified in the register. Third-party publication and data rights remain with their owners. There are no private patient-level records in this package. Any institutional ethics determination required for future patient or developer research remains to be obtained.')
p=para('Repository: ');hyperlink(p,'clinical-trials/medicare','https://github.com/clinical-trials/medicare')

# Only substantive citations from the manuscript, tables and supplements are printed.
missing=[rid for rid in used if rid not in refs]
assert not missing,f'Unresolved citation IDs: {missing}'
page();heading('References')
para('Stable IDs support the working manuscript until final journal numbering. Numeric PMIDs identify scientific and methodological publications; policy/data sources have no PMID. The separate living bibliography and RIS retain all active, excluded, contextual and correction records. Bibliographic verification is distinct from eligibility and full-text review.')
for rid in used:
    r=refs[rid]
    if rid.startswith('P-'):
        title=r.get('official_title') or '; '.join(r.get('descriptions',[]))
        org=r.get('organization')
        value=f"[{rid}] "+(f"{org}. " if org else '')+f"{title}."
        if not r.get('official_title'):value+=' Descriptive source label.'
        value+=' Access/status: '+str(r.get('source_access_date') or r.get('recorded_on') or 'See source register')+'.'
        url=r['url']
    else:
        value=f"[{rid}] {r.get('citation') or r.get('title')}"
        if r.get('pmid'):value+=f" PMID {r['pmid']}."
        else:value+=' PMID not available; see lookup status in source registry.'
        if r.get('doi'):value+=f" DOI {r['doi']}."
        if rid=='L-B01-E1':value+=' Correction linked to L-B01; not an independent study.'
        if rid=='L-T03-P1':value+=' Same trial as L-T03; not an independent trial.'
        url=r.get('source_url') or r.get('url')
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.0
    p.paragraph_format.keep_together=True
    run=p.add_run(value);run.font.size=Pt(10)
    if url:
        p.add_run(' ');hyperlink(p,'Source',url)

OUT.mkdir(parents=True,exist_ok=True)
path=OUT/(STEM+'.docx');doc.save(path)
stats={'built_on':'2026-09-21','title_characters':len(source.splitlines()[0][2:]),
 'abstract_words_whitespace_count':len(source.split('## Abstract')[1].split('## Key takeaways')[0].split()),
 'main_body_words_including_headings_and_stable_ids':len(source.split('## Introduction')[1].split()),
 'main_figures':2,'main_tables':2,'opportunity_domains':6,'printed_reference_count':len(used),
 'printed_reference_ids':used,'docx_path':str(path),'formal_human_review_completed':False,
 'word_count_note':'Working count; journal treatment of headings, citations, tables and supplements must be confirmed.'}
(HERE/'publication_build.json').write_text(json.dumps(stats,indent=2)+'\n')
print(json.dumps(stats,indent=2))
